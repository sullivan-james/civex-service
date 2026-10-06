"""Putting a table on disk, in any of `civex.domain.tables.FORMATS`, one page of
rows at a time so a large export never holds its rows in memory."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Iterable

from civex.domain import tables

# openpyxl refuses these control characters in a cell.
_ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def write_table(
    dest: Path,
    fmt: str,
    columns: list[str],
    batches: Iterable[list[dict[str, Any]]],
    sheet: str = "Sheet",
) -> int:
    """Write the rows to `dest` in `fmt`; returns how many rows were written. A
    row is a dict keyed by column; a value it lacks is blank."""
    tables.check_format(fmt)
    if fmt in tables.TEXT_FORMATS:
        return _write_delimited(dest, fmt, columns, batches)
    if fmt == "xlsx":
        return _write_xlsx(dest, columns, batches, sheet)
    return _write_json(dest, fmt, batches)


def _write_delimited(
    dest: Path, fmt: str, columns: list[str], batches: Iterable[list[dict[str, Any]]]
) -> int:
    count = 0
    with dest.open("w", encoding="utf-8", newline="") as out:
        writer = csv.writer(out, delimiter="\t" if fmt == "tsv" else ",")
        writer.writerow(columns)
        for rows in batches:
            for row in rows:
                writer.writerow([tables.cell(row.get(c), fmt) for c in columns])
                count += 1
    return count


def _write_json(dest: Path, fmt: str, batches: Iterable[list[dict[str, Any]]]) -> int:
    """`json`: one array (2-space indent); `jsonl`: one object per line."""
    count = 0
    with dest.open("w", encoding="utf-8", newline="") as out:
        for rows in batches:
            for row in rows:
                body = json.dumps(
                    tables.nest(row), default=str, indent=None if fmt == "jsonl" else 2
                )
                if fmt == "jsonl":
                    out.write(body + "\n")
                else:
                    out.write("[\n" if count == 0 else ",\n")
                    out.write("  " + body.replace("\n", "\n  "))
                count += 1
        if fmt == "json":
            out.write("[]" if count == 0 else "\n]")
    return count


def _write_xlsx(
    dest: Path,
    columns: list[str],
    batches: Iterable[list[dict[str, Any]]],
    sheet: str,
) -> int:
    from openpyxl import Workbook

    book = Workbook(write_only=True)
    ws = book.create_sheet(tables.sheet_name(sheet))
    ws.freeze_panes = "A2"
    ws.append(columns)
    count = 0
    for rows in batches:
        for row in rows:
            ws.append([_xlsx_value(row.get(c)) for c in columns])
            count += 1
    book.save(dest)
    return count


def _xlsx_value(value: Any) -> Any:
    value = tables.cell(value, "xlsx")
    if isinstance(value, str):
        return _ILLEGAL_XLSX.sub("", value)
    if value is None or isinstance(value, (bool, int, float)):
        return value
    return str(value)
