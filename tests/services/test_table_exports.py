"""Tables beside (or instead of) files in an export: the folder, the zip, and
the saved definition that asks for them."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from civex.context import AppContext
from civex.domain.exceptions import ValidationError
from civex.domain.file_access import FileSelection
from civex.domain.query import RecordQuery
from civex.domain.tables import TableSpec
from civex.services.archive import build_download
from civex.services.file_access_service import GENERATED, MARKER


@pytest.fixture()
def study(ctx: AppContext, make_schema, make_collection, make_record):
    """An encounter with two recordings, each with a table file."""
    make_schema("encounter", fields=[("name", "string")])
    make_schema(
        "recording",
        fields=[("rname", "string"), ("depth", "float"), ("table", "file")],
        parent="encounter",
    )
    make_collection("hb")
    enc = make_record("hb", "encounter", {"name": "Encounter 7"})
    for name, depth in (("A", 1.5), ("B", 2.5)):
        ref = ctx.file_svc.store_bytes(name.encode(), f"{name}.txt").to_dict()
        make_record(
            "hb",
            "recording",
            {"rname": name, "depth": depth, "table": ref},
            parent_record_id=str(enc.id),
        )
    # A recording with no file at all: still a row in the table.
    make_record(
        "hb", "recording", {"rname": "C", "depth": 3.5}, parent_record_id=str(enc.id)
    )
    return enc


def _sel(**kw) -> FileSelection:
    table = kw.pop("table", TableSpec("csv"))
    files = kw.pop("files", True)
    layout = kw.pop("layout", "tree")
    return FileSelection(
        query=RecordQuery(schema="recording", **kw),
        tables=[table] if table else [],
        files=files,
        layout=layout,
    )


def _rows(path: Path) -> list[dict]:
    with path.open(newline="") as f:
        return list(csv.DictReader(f))


def test_the_table_lists_the_records_and_says_where_each_file_is(
    ctx: AppContext, study, tmp_path: Path
) -> None:
    sel = _sel(table=TableSpec("csv", ["rname", "depth", "table"]))
    dest = tmp_path / "out"

    result = ctx.file_access_svc.export(sel, dest, "copy")

    assert result.tables == 1
    rows = _rows(dest / "Recordings.csv")
    assert rows == [
        {"rname": "A", "depth": "1.5", "table": "Encounter 7/A/A.txt"},
        {"rname": "B", "depth": "2.5", "table": "Encounter 7/B/B.txt"},
        {"rname": "C", "depth": "3.5", "table": ""},
    ]
    for row in rows[:2]:
        assert (dest / row["table"]).is_file()  # the cell is a path in the folder


def test_the_default_table_has_the_id_and_every_field_own_and_inherited(
    ctx: AppContext, study, tmp_path: Path
) -> None:
    plan = ctx.file_access_svc.plan(_sel())

    (table,) = plan.tables
    assert table.columns == ["id", "rname", "depth", "table", "name"]
    assert (table.name, table.kind, table.rows) == ("Recordings.csv", "recording", 3)


def test_columns_the_kind_does_not_have_are_left_out(ctx: AppContext, study) -> None:
    plan = ctx.file_access_svc.plan(
        _sel(table=TableSpec("csv", ["rname", "nonsense", "updated_at"]))
    )

    assert plan.tables[0].columns == ["rname", "updated_at"]


def test_a_table_is_re_made_and_dropped_when_no_longer_asked_for(
    ctx: AppContext, study, tmp_path: Path
) -> None:
    dest = tmp_path / "out"
    ctx.file_access_svc.export(_sel(), dest, "copy")
    marker = json.loads((dest / MARKER).read_text())["files"]
    assert marker["Recordings.csv"]["how"] == GENERATED

    ctx.file_access_svc.export(_sel(table=TableSpec("tsv")), dest, "copy")
    assert (dest / "Recordings.tsv").is_file()
    assert not (dest / "Recordings.csv").exists()  # the old one is pruned

    result = ctx.file_access_svc.export(_sel(table=None), dest, "copy")
    assert not (dest / "Recordings.tsv").exists()
    assert result.removed == 1 and result.tables == 0


def test_a_table_alone_needs_no_files_and_goes_in_the_project(
    ctx: AppContext, study
) -> None:
    sel = _sel(files=False, table=TableSpec("jsonl", ["rname"]))

    plan = ctx.file_access_svc.plan(sel)
    result = ctx.file_access_svc.export_managed(sel, "just-the-table", "link")

    assert plan.items == [] and plan.complete
    out = Path(result.dest)
    assert [
        json.loads(x) for x in (out / "Recordings.jsonl").read_text().split("\n") if x
    ] == [
        {"rname": "A"},
        {"rname": "B"},
        {"rname": "C"},
    ]
    assert result.location == "project"


def test_beside_files_only_kinds_that_hold_files_get_a_table(
    ctx: AppContext, study
) -> None:
    sel = FileSelection(
        query=RecordQuery(within=str(study.id)),
        tables=[TableSpec("csv")],
        files=True,
    )
    alone = FileSelection(
        query=RecordQuery(within=str(study.id)),
        tables=[TableSpec("csv")],
        files=False,
    )

    beside = ctx.file_access_svc.plan(sel)
    on_its_own = ctx.file_access_svc.plan(alone)

    assert [t.kind for t in beside.tables] == ["recording"]
    assert [t.kind for t in on_its_own.tables] == ["encounter", "recording"]


def test_a_table_never_takes_the_name_of_a_file_and_never_writes_through_a_link(
    ctx: AppContext, study, tmp_path: Path
) -> None:
    # A flat layout puts every file at the top: here one is called like the table.
    ref = ctx.file_svc.store_bytes(b"precious", "Recordings.csv").to_dict()
    ctx.record_svc.add(
        "hb",
        "recording",
        {"rname": "D", "table": ref},
        parent_record_id=str(study.id),
    )
    sel = _sel(layout="flat", table=TableSpec("csv", ["rname"]))
    dest = tmp_path / "out"

    plan = ctx.file_access_svc.plan(sel)
    ctx.file_access_svc.export(sel, dest, "link")

    assert plan.tables[0].name == "Recordings (2).csv"
    assert (dest / "Recordings.csv").read_bytes() == b"precious"
    assert _rows(dest / "Recordings (2).csv")[-1] == {"rname": "D"}


def test_download_is_a_zip_of_files_and_table_or_the_lone_table(
    ctx: AppContext, study, tmp_path: Path
) -> None:
    svc = ctx.file_access_svc
    both = _sel(table=TableSpec("xlsx", ["rname", "table"]))
    scratch = tmp_path / "s1"
    scratch.mkdir()
    zipped = build_download(svc, ctx.file_svc, both, svc.plan(both), scratch, "mine")

    assert zipped.filename == "mine.zip"
    zf = zipfile.ZipFile(zipped.path)
    assert sorted(zf.namelist()) == [
        "Encounter 7/A/A.txt",
        "Encounter 7/B/B.txt",
        "Recordings.xlsx",
    ]
    book = load_workbook(io.BytesIO(zf.read("Recordings.xlsx")))
    assert [c.value for c in book.active["B"]] == [
        "table",
        "Encounter 7/A/A.txt",
        "Encounter 7/B/B.txt",
        None,
    ]

    alone = _sel(files=False, table=TableSpec("tsv", ["rname"]))
    scratch = tmp_path / "s2"
    scratch.mkdir()
    lone = build_download(svc, ctx.file_svc, alone, svc.plan(alone), scratch, "names")
    assert lone.filename == "names.tsv"
    assert lone.path.read_text().splitlines() == ["rname", "A", "B", "C"]


def test_there_is_nothing_to_download_when_nothing_is_selected(
    ctx: AppContext, study, tmp_path: Path
) -> None:
    empty = _sel(files=False, table=None)
    with pytest.raises(ValidationError):
        build_download(
            ctx.file_access_svc,
            ctx.file_svc,
            empty,
            ctx.file_access_svc.plan(empty),
            tmp_path,
        )


def test_a_format_must_be_one_that_exists() -> None:
    with pytest.raises(ValidationError):
        TableSpec("parquet")


# --- saved with a schema ---


def test_a_saved_export_can_be_a_table_beside_files_or_a_table_alone(
    ctx: AppContext, study
) -> None:
    svc = ctx.export_def_svc
    both = svc.create(
        "encounter",
        "Tables and files",
        holder="recording",
        tables=[{"format": "xlsx", "columns": ["rname", "depth"]}],
    )
    alone = svc.create(
        "encounter",
        "Just the table",
        holder="recording",
        include_files=False,
        tables=[{"format": "csv", "columns": None}],
    )

    run = svc.selection(both, collection="hb")
    only = svc.selection(alone, collection="hb")

    assert run.files and run.tables == [TableSpec("xlsx", ["rname", "depth"])]
    assert not only.files and only.tables == [TableSpec("csv", None)]
    assert only.fields is None


def test_a_saved_export_needs_files_a_table_or_both(ctx: AppContext, study) -> None:
    with pytest.raises(ValidationError, match="files, tables, or both"):
        ctx.export_def_svc.create(
            "encounter", "Nothing", holder="recording", include_files=False
        )


def test_a_saved_table_is_checked_against_the_kind_it_lists(
    ctx: AppContext, study
) -> None:
    with pytest.raises(ValidationError, match="'nonsense' isn't a field"):
        ctx.export_def_svc.create(
            "encounter",
            "Bad column",
            holder="recording",
            tables=[{"format": "csv", "columns": ["nonsense"]}],
        )
    with pytest.raises(ValidationError, match="format is one of"):
        ctx.export_def_svc.create(
            "encounter", "Bad format", holder="recording", tables=[{"format": "pdf"}]
        )


def test_turning_files_off_keeps_a_table_and_clears_the_file_fields(
    ctx: AppContext, study
) -> None:
    svc = ctx.export_def_svc
    svc.create(
        "encounter",
        "Both",
        holder="recording",
        fields=["table"],
        tables=[{"format": "csv", "columns": None}],
    )

    changed = svc.update("encounter", "Both", include_files=False)

    assert changed.include_files is False and changed.fields == []
    assert changed.tables == [{"format": "csv", "columns": None}]
