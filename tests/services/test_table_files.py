"""Writing a table in each format, and what a cell holds in it."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from civex.domain import tables
from civex.domain.exceptions import ValidationError
from civex.services.table_files import write_table

POINT = {"type": "Point", "coordinates": [-3.4, 56.1]}
ROWS = [
    {"name": "A", "n": 1, "tags": ["x", "y"], "where": POINT, "ok": True},
    {"name": "B", "n": None, "tags": [], "where": None, "ok": False},
]
COLUMNS = ["name", "n", "tags", "where", "ok"]


def _write(tmp_path: Path, fmt: str, batches=None) -> Path:
    dest = tmp_path / f"out{tables.extension(fmt)}"
    write_table(dest, fmt, COLUMNS, batches or [ROWS[:1], ROWS[1:]])
    return dest


@pytest.mark.parametrize("fmt,delimiter", [("csv", ","), ("tsv", "\t")])
def test_text_tables_write_lists_as_json_and_points_as_text(
    tmp_path: Path, fmt: str, delimiter: str
) -> None:
    dest = _write(tmp_path, fmt)

    with dest.open(newline="") as f:
        rows = list(csv.reader(f, delimiter=delimiter))
    assert rows[0] == COLUMNS
    assert rows[1] == ["A", "1", '["x", "y"]', "56.1, -3.4", "True"]
    assert rows[2] == ["B", "", "[]", "", "False"]


def test_json_keeps_values_as_they_are_and_nests_joined_columns(
    tmp_path: Path,
) -> None:
    dest = tmp_path / "out.json"
    write_table(
        dest, "json", ["a", "c.d"], [[{"a": 1, "c.d": "x"}], [{"a": 2, "c.d": None}]]
    )

    assert json.loads(dest.read_text()) == [
        {"a": 1, "c": {"d": "x"}},
        {"a": 2, "c": {"d": None}},
    ]


def test_an_empty_table_is_still_valid(tmp_path: Path) -> None:
    for fmt, empty in (("json", []), ("jsonl", None), ("csv", None)):
        dest = tmp_path / f"e{tables.extension(fmt)}"
        assert write_table(dest, fmt, ["a"], []) == 0
        if fmt == "json":
            assert json.loads(dest.read_text()) == empty
        elif fmt == "jsonl":
            assert dest.read_text() == ""
        else:
            assert dest.read_text().strip() == "a"


def test_jsonl_is_one_object_per_line(tmp_path: Path) -> None:
    dest = _write(tmp_path, "jsonl")

    lines = dest.read_text().splitlines()
    assert [json.loads(line)["name"] for line in lines] == ["A", "B"]
    assert json.loads(lines[0])["tags"] == ["x", "y"]  # native, not text


def test_xlsx_keeps_numbers_and_booleans_and_names_the_sheet(tmp_path: Path) -> None:
    dest = tmp_path / "out.xlsx"
    write_table(dest, "xlsx", COLUMNS, [ROWS], sheet="Sel: ections/1")

    book = load_workbook(dest)
    sheet = book.active
    assert sheet.title == "Sel  ections 1"  # characters Excel refuses are replaced
    assert [c.value for c in sheet[1]] == COLUMNS
    assert [c.value for c in sheet[2]] == ["A", 1, '["x", "y"]', "56.1, -3.4", True]


def test_xlsx_drops_control_characters_excel_refuses(tmp_path: Path) -> None:
    dest = tmp_path / "out.xlsx"
    write_table(dest, "xlsx", ["a"], [[{"a": "bad\x00text"}]])

    assert load_workbook(dest).active["A2"].value == "badtext"


def test_an_unknown_format_is_refused(tmp_path: Path) -> None:
    with pytest.raises(ValidationError):
        write_table(tmp_path / "x", "parquet", ["a"], [])
